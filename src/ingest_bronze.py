"""Bronze: landing Excel -> Parquet, 1:1. All source columns kept as strings; no cleaning or filtering."""

from pathlib import Path

import openpyxl
import pandas as pd

from src import config
from src.utils import get_logger, ist_now, list_landing_files, new_batch_id, quarantine_file

log = get_logger(__name__)


class FileQuarantinedError(RuntimeError):
    """Raised after a run in which at least one file was quarantined (plan section 9, step 6)."""


# --------------------------------------------------------------------------- #
# Layout detection and expected row counts
# --------------------------------------------------------------------------- #
def detect_layout(path: Path):
    """Return the layout key ('A' or 'B') by sheet names, or None if no layout matches."""
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        sheets = set(wb.sheetnames)
    finally:
        wb.close()
    for key, layout in config.SOURCE_LAYOUTS.items():
        if layout["match_sheets"] <= sheets:
            return key
    return None


def count_excel_data_rows(path: Path, layout_key: str) -> int:
    """Independent count of data rows in the data sheet (header rows excluded,
    footer rows included). Used by validate_bronze to check nothing was lost."""
    layout = config.SOURCE_LAYOUTS[layout_key]
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        last_non_empty = 0
        for i, row in enumerate(wb[layout["data_sheet"]].iter_rows(values_only=True), start=1):
            if any(v is not None for v in row):
                last_non_empty = i
    finally:
        wb.close()
    return max(last_non_empty - layout["header_rows"], 0)


# --------------------------------------------------------------------------- #
# Reading
# --------------------------------------------------------------------------- #
def build_column_names(header: pd.DataFrame) -> list:
    """One name per column from the header rows. The lowest header row with a
    value wins (Layout A: row 2 display name, falling back to row 1), then the
    name is sanitised. Empty -> col_<n>. Duplicates get a _<n> suffix."""
    names, seen = [], set()
    for j in range(header.shape[1]):
        label = None
        for i in reversed(range(header.shape[0])):
            value = header.iat[i, j]
            if pd.notna(value) and str(value).strip():
                label = str(value)
                break
        name = config.sanitize_column_name(label) if label else ""
        if not name:
            name = f"col_{j + 1}"
        if name in seen:
            name = f"{name}_{j + 1}"
        seen.add(name)
        names.append(name)
    return names


def read_excel_raw(path: Path, layout_key: str) -> pd.DataFrame:
    layout = config.SOURCE_LAYOUTS[layout_key]
    raw = pd.read_excel(
        path,
        sheet_name=layout["data_sheet"],
        header=None,
        dtype=str,
        keep_default_na=False,   # keep text like "#N/A" or "NA" as-is
        na_values=[""],          # only truly empty cells become null
    )
    n_header = layout["header_rows"]
    df = raw.iloc[n_header:].reset_index(drop=True)
    df.columns = build_column_names(raw.iloc[:n_header])
    return df.astype("string")


# --------------------------------------------------------------------------- #
# One file
# --------------------------------------------------------------------------- #
def ingest_file(path: Path, layout_key: str, batch_id: str, ingested_at) -> dict:
    layout = config.SOURCE_LAYOUTS[layout_key]
    df = read_excel_raw(path, layout_key)

    missing = [c for c in (layout["timestamp_col"], layout["demand_col"]) if c not in df.columns]
    if missing:
        raise ValueError(f"missing_columns: {missing}")

    df["_source_file"] = pd.array([path.name] * len(df), dtype="string")
    df["_source_layout"] = pd.array([layout_key] * len(df), dtype="string")
    df["_ingested_at"] = pd.Timestamp(ingested_at)
    df["_batch_id"] = pd.array([batch_id] * len(df), dtype="string")

    out_dir = config.BRONZE_DIR / f"ingest_date={ingested_at:%Y-%m-%d}"
    out_dir.mkdir(parents=True, exist_ok=True)
    final = out_dir / f"{path.stem}.parquet"      # same-day rerun overwrites the same file
    tmp = out_dir / f".{path.stem}.parquet.tmp"   # hidden name: readers ignore it
    df.to_parquet(tmp, engine="pyarrow", index=False)
    tmp.replace(final)

    return {
        "file": path.name,
        "layout": layout_key,
        "rows_written": len(df),
        "parquet": str(final),
    }


# --------------------------------------------------------------------------- #
# Whole landing folder
# --------------------------------------------------------------------------- #
def run(batch_id: str = None) -> dict:
    batch_id = batch_id or new_batch_id()
    ingested_at = ist_now()
    files = list_landing_files()
    log.info("bronze start | batch=%s | files in landing=%d", batch_id, len(files))

    ingested, quarantined = [], []
    for path in files:
        try:
            layout_key = detect_layout(path)
        except Exception as exc:  # corrupt zip, not an xlsx, password-protected, ...
            log.error("unreadable file %s: %s", path.name, exc)
            quarantine_file(path, f"{config.REASON_UNREADABLE_FILE}: {exc}")
            quarantined.append({"file": path.name, "reason": config.REASON_UNREADABLE_FILE})
            continue

        if layout_key is None:
            log.error("unknown layout: %s", path.name)
            quarantine_file(path, config.REASON_UNKNOWN_LAYOUT)
            quarantined.append({"file": path.name, "reason": config.REASON_UNKNOWN_LAYOUT})
            continue

        try:
            result = ingest_file(path, layout_key, batch_id, ingested_at)
        except ValueError as exc:
            log.error("cannot ingest %s: %s", path.name, exc)
            quarantine_file(path, str(exc))
            quarantined.append({"file": path.name, "reason": config.REASON_MISSING_COLUMNS})
            continue

        result["expected_rows"] = count_excel_data_rows(path, layout_key)
        log.info(
            "bronze ok | %s | layout=%s | rows in=%d out=%d",
            path.name, layout_key, result["expected_rows"], result["rows_written"],
        )
        ingested.append(result)

    summary = {"batch_id": batch_id, "ingested": ingested, "quarantined": quarantined}
    log.info("bronze done | ingested=%d quarantined=%d", len(ingested), len(quarantined))

    if quarantined:
        raise FileQuarantinedError(f"{len(quarantined)} file(s) quarantined: {quarantined}")
    return summary